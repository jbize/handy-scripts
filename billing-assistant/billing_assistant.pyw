#!/usr/bin/env python3
"""
Billing Assistant
Version: 3.4.0
"""

from __future__ import annotations

import calendar
import configparser
import errno
import logging
import os
import re
import shutil
import sys
import webbrowser
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote

import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext

try:
    import openpyxl
except ImportError:
    openpyxl = None

# --- CONSTANTS ---
APP_NAME = "Billing Assistant"
APP_VERSION = "3.4.0"
CONFIG_FILE = "config.ini"
LOG_FILE = "billing_assistant.log"

# Trailing _YYYY-MM on a workbook stem (e.g. "Time Tracking_2026-08")
WORKBOOK_MONTH_SUFFIX = re.compile(r"_\d{4}-\d{2}$")

# mailto: bodies often strip leading ASCII spaces; NBSP usually survives.
EMAIL_INDENT = "\u00a0" * 4

# Workbook layout (1-based openpyxl coordinates)
CELL_RATE = (1, 5)       # E1
CELL_HOURS = (15, 2)     # B15
CELL_AMOUNT = (18, 2)    # B18
CELL_MONTH = (5, 4)      # D5
DATE_COL = 4
DATE_FIRST_ROW = 5
DATE_SLOT_COUNT = 31
TIME_ENTRY_COLS = (7, 8, 10, 11, 13, 14)  # G H J K M N
TIME_ENTRY_ROWS = range(5, 39)

# Windows sharing / lock codes (Excel or OneDrive holding the file)
_WIN_SHARING_VIOLATION = 32
_WIN_LOCK_VIOLATION = 33

WINDOW_COLLAPSED = "750x220"
WINDOW_EXPANDED = "750x750"


# --- DATA CLASSES ---
@dataclass(slots=True)
class InvoiceEmail:
    to: str
    cc: str
    subject: str
    body: str


@dataclass(slots=True)
class BillingData:
    rate: float
    hours: float
    amount: float
    month_label: str
    config: dict[str, str]


# --- CONFIGURATION MANAGER ---
class ConfigManager:
    def __init__(self) -> None:
        self.app_dir = Path(os.environ.get("APPDATA", "~")).expanduser() / APP_NAME
        self.app_dir.mkdir(parents=True, exist_ok=True)
        self.path = self.app_dir / CONFIG_FILE
        self.log_path = self.app_dir / LOG_FILE
        self.config = configparser.ConfigParser()
        self.load()

    def load(self) -> None:
        if self.path.exists():
            self.config.read(self.path, encoding="utf-8")
        if "Settings" not in self.config:
            self.config["Settings"] = {
                "last_directory": str(Path.home()),
                "last_file": "",
                "remember_file": "True",
                "enable_logging": "False",
            }

    def save(self) -> None:
        with open(self.path, "w", encoding="utf-8") as f:
            self.config.write(f)

    def get(self, key: str, fallback: str = "") -> str:
        return self.config["Settings"].get(key, fallback)

    def set(self, key: str, value: str) -> None:
        self.config["Settings"][key] = value
        self.save()

    def update(self, **values: str) -> None:
        for key, value in values.items():
            self.config["Settings"][key] = value
        self.save()

    def get_bool(self, key: str, fallback: bool = False) -> bool:
        return self.config["Settings"].getboolean(key, fallback=fallback)


# --- LOGGING / ERRORS ---
def setup_logging(cfg: ConfigManager) -> None:
    root_logger = logging.getLogger()
    root_logger.handlers.clear()
    if cfg.get_bool("enable_logging"):
        logging.disable(logging.NOTSET)
        logging.basicConfig(
            filename=cfg.log_path,
            level=logging.INFO,
            format="%(asctime)s - %(levelname)s - %(message)s",
        )
    else:
        logging.disable(logging.CRITICAL)


def is_permission_denied(exc: BaseException) -> bool:
    if isinstance(exc, PermissionError):
        return True
    if not isinstance(exc, OSError):
        return False
    if exc.errno in (errno.EACCES, errno.EPERM):
        return True
    return getattr(exc, "winerror", None) in (
        _WIN_SHARING_VIOLATION,
        _WIN_LOCK_VIOLATION,
    )


def workbook_locked_message(path: Path) -> str:
    return (
        f"Cannot read:\n{path}\n\n"
        "The file appears to be open in Excel or locked by Windows/OneDrive.\n"
        "Please close the workbook in Excel, let OneDrive finish syncing if needed, and try again."
    )


def open_workbook(path: Path, **kwargs: Any):
    """Load a workbook, mapping lock/sharing failures to a clear PermissionError."""
    try:
        return openpyxl.load_workbook(path, **kwargs)
    except Exception as e:
        if is_permission_denied(e):
            raise PermissionError(workbook_locked_message(path)) from e
        raise


def normalize_config_key(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_")


def workbook_basename(path: Path) -> str:
    """Basename for rolled files: stem with trailing _YYYY-MM removed."""
    return WORKBOOK_MONTH_SUFFIX.sub("", path.stem)


def rolled_workbook_path(path: Path, new_start: datetime) -> Path:
    """Target path beside the source, preserving the workbook basename."""
    return path.parent / f"{workbook_basename(path)}_{new_start.strftime('%Y-%m')}{path.suffix}"


def cell_float(ws, row: int, column: int) -> float:
    value = ws.cell(row=row, column=column).value
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError) as e:
        coord = ws.cell(row=row, column=column).coordinate
        raise ValueError(f"Cell {coord} must be a number (got {value!r}).") from e


# --- EXCEL PARSING ---
def read_config_block(ws) -> dict[str, str]:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value != "__CONFIG__":
                continue
            config: dict[str, str] = {}
            check_row = cell.row + 1
            col_idx = cell.column
            while True:
                key = ws.cell(row=check_row, column=col_idx).value
                if not key:
                    return config
                value = ws.cell(row=check_row, column=col_idx + 1).value
                config[normalize_config_key(key)] = "" if value is None else str(value).strip()
                check_row += 1
    raise ValueError("__CONFIG__ marker not found in the workbook.")


def parse_workbook(path: Path) -> BillingData:
    wb = open_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        config = read_config_block(ws)
        rate = cell_float(ws, *CELL_RATE)
        hours = cell_float(ws, *CELL_HOURS)
        amount = cell_float(ws, *CELL_AMOUNT)
        d5 = ws.cell(row=CELL_MONTH[0], column=CELL_MONTH[1]).value
        month_label = d5.strftime("%B-%Y") if isinstance(d5, datetime) else "Unknown"
        return BillingData(rate, hours, amount, month_label, config)
    finally:
        wb.close()


# --- PAYMENT / EMAIL ---
def indent_lines(text: str, prefix: str = EMAIL_INDENT) -> str:
    return "\n".join(f"{prefix}{line}" if line else line for line in text.splitlines())


def build_payment_block(config: dict[str, str], amount: float) -> str:
    method_name = str(config.get("current_payment", "")).strip().lower()
    if not method_name:
        return ""

    if method_name == "paypal":
        parts: list[str] = []
        paypal_email = config.get("paypal_email", "")
        paypal_me = config.get("paypal_me", "")
        if paypal_email:
            parts.append(
                "Please submit payment via PayPal Friends & Family to:\n"
                + indent_lines(paypal_email)
            )
        if paypal_me:
            parts.append(
                "Or use PayPal.Me:\n"
                + indent_lines(f"https://paypal.me/{paypal_me}/{amount:.2f}")
            )
        return "\n\n".join(parts)

    if method_name == "zelle":
        zelle_address = config.get("zelle_address", "")
        if not zelle_address:
            return "Please submit payment via Zelle."
        return "Please submit payment via Zelle to:\n" + indent_lines(zelle_address)

    if method_name == "check":
        address = config.get("check_address", "")
        if not address:
            return "Please send payment by check."
        return "Please send check to:\n" + indent_lines(address)

    # Custom: current_payment names another config key whose value is free-form text.
    custom_instr = config.get(method_name, "")
    if custom_instr:
        return str(custom_instr).strip()
    return ""


def build_email(data: BillingData) -> InvoiceEmail:
    config = data.config
    slug = config.get("invoice_slug", "invoice")
    sender = config.get("sender_first_name", "Consultant")
    subject = f"{sender}-{slug} ({data.month_label})"
    to = config.get("client_email", "")
    # CC priority: cc_email -> sender_email -> paypal_email
    cc = config.get("cc_email") or config.get("sender_email") or config.get("paypal_email") or ""

    body_parts = [
        f"{config.get('client_greeting', 'Client')},",
        "",
        f"{data.month_label} {slug}:",
        "",
        f"{EMAIL_INDENT}• Hours worked: {round(data.hours, 2)} @ ${data.rate:.0f}/hr",
        f"{EMAIL_INDENT}• Amount due:   ${data.amount:.2f}",
    ]
    payment_block = build_payment_block(config, data.amount)
    if payment_block:
        body_parts.extend(["", payment_block])
    body_parts.extend(["", f"Thanks,\n{sender}"])
    return InvoiceEmail(to, cc, subject, "\n".join(body_parts))


# --- MONTH ROLL ---
def month_start(dt: datetime) -> datetime:
    return datetime(dt.year, dt.month, 1)


def next_month_start(dt: datetime) -> datetime:
    return datetime(
        dt.year + (dt.month // 12),
        (dt.month % 12) + 1,
        1,
    )


def read_sheet_month(path: Path) -> datetime:
    """Return the first-of-month date from cell D5."""
    wb = open_workbook(path, data_only=True, read_only=True)
    try:
        curr_date = wb.active.cell(row=CELL_MONTH[0], column=CELL_MONTH[1]).value
    finally:
        wb.close()

    if not isinstance(curr_date, datetime):
        raise ValueError("Cell D5 must contain a valid date.")
    return month_start(curr_date)


def apply_month_template(ws, new_start: datetime) -> None:
    """Rewrite the date column for new_start and clear time-entry cells."""
    days_in_month = calendar.monthrange(new_start.year, new_start.month)[1]
    for i in range(DATE_SLOT_COUNT):
        cell = ws.cell(row=DATE_FIRST_ROW + i, column=DATE_COL)
        cell.value = (new_start + timedelta(days=i)) if i < days_in_month else None

    for row_idx in TIME_ENTRY_ROWS:
        for col_idx in TIME_ENTRY_COLS:
            ws.cell(row=row_idx, column=col_idx).value = None


def roll_month(path: Path, target_start: datetime | None = None) -> Path:
    """
    Copy the workbook to target_start (default: month after D5), rewrite dates,
    and clear time-entry cells.
    """
    sheet_month = read_sheet_month(path)
    new_start = month_start(target_start) if target_start else next_month_start(sheet_month)
    new_path = rolled_workbook_path(path, new_start)

    if new_path.resolve() == path.resolve():
        raise ValueError(
            f"Roll target is the same file:\n{new_path}\n"
            "Choose a different month or rename the source workbook first."
        )

    shutil.copy2(path, new_path)
    try:
        wb = open_workbook(new_path, data_only=False)
        try:
            apply_month_template(wb.active, new_start)
            try:
                wb.save(new_path)
            except Exception as e:
                if is_permission_denied(e):
                    raise PermissionError(workbook_locked_message(new_path)) from e
                raise
        finally:
            wb.close()
    except Exception:
        if new_path.exists() and new_path != path:
            try:
                new_path.unlink()
            except OSError:
                logging.warning("Could not remove incomplete roll file: %s", new_path)
        raise

    return new_path


# --- DIALOGS ---
def show_error(parent: tk.Tk, message: str) -> None:
    dlg = tk.Toplevel(parent)
    dlg.title("Error")
    dlg.geometry("560x320")
    dlg.transient(parent)
    dlg.grab_set()
    tk.Label(dlg, text="Error", fg="red", font=("", 10, "bold")).pack(pady=8)
    txt = scrolledtext.ScrolledText(dlg, wrap="word", padx=10, pady=10)
    txt.insert("1.0", message)
    txt.focus_set()
    txt.pack(fill="both", expand=True, padx=15, pady=5)
    tk.Button(dlg, text="Close", command=dlg.destroy, width=15).pack(pady=8)
    dlg.protocol("WM_DELETE_WINDOW", dlg.destroy)


def ask_catch_up_roll(
    parent: tk.Tk,
    sheet_month: datetime,
    next_start: datetime,
    current_start: datetime,
) -> datetime | None:
    """
    When the timesheet month is behind the calendar, ask whether to jump to
    the current month or roll only one month forward. Returns the chosen
    month-start date, or None if cancelled.
    """
    result: dict[str, datetime | None] = {"value": None}
    sheet_label = sheet_month.strftime("%B %Y")
    next_label = next_start.strftime("%B %Y")
    current_label = current_start.strftime("%B %Y")

    dlg = tk.Toplevel(parent)
    dlg.title("Roll timesheet")
    dlg.geometry("520x210")
    dlg.transient(parent)
    dlg.grab_set()

    msg = (
        f"This timesheet is for {sheet_label}.\n\n"
        f"Create {current_label} (current month), or roll to {next_label}?"
    )
    tk.Label(dlg, text=msg, justify="left", padx=16, pady=16).pack(anchor="w")

    btn_row = tk.Frame(dlg)
    btn_row.pack(pady=8)

    def choose(value: datetime | None) -> None:
        result["value"] = value
        dlg.destroy()

    tk.Button(
        btn_row,
        text=f"Current month ({current_label})",
        width=28,
        command=lambda: choose(current_start),
    ).pack(side="left", padx=5)
    tk.Button(
        btn_row,
        text=f"Next month ({next_label})",
        width=24,
        command=lambda: choose(next_start),
    ).pack(side="left", padx=5)
    tk.Button(btn_row, text="Cancel", width=10, command=lambda: choose(None)).pack(
        side="left", padx=5
    )

    dlg.protocol("WM_DELETE_WINDOW", lambda: choose(None))
    parent.wait_window(dlg)
    return result["value"]


# --- GUI ---
class BillingAssistantGUI:
    def __init__(self, root: tk.Tk, args: list[str]) -> None:
        self.root = root
        self.root.title(f"{APP_NAME} {APP_VERSION}")
        self.root.geometry(WINDOW_COLLAPSED)

        self.cfg = ConfigManager()
        setup_logging(self.cfg)
        self._setup_ui()

        if len(args) > 1:
            self._set_file(Path(args[1]))
        elif self.cfg.get_bool("remember_file"):
            last = self.cfg.get("last_file")
            if last and Path(last).exists():
                self._set_file(Path(last))

    def _setup_ui(self) -> None:
        pad = {"padx": 15, "pady": 5}

        f_frame = tk.LabelFrame(self.root, text="Timesheet Source", padx=10, pady=10)
        f_frame.pack(fill="x", **pad)

        tk.Label(f_frame, text="Directory:").grid(row=0, column=0, sticky="e")
        self.dir_var = tk.StringVar(value=self.cfg.get("last_directory"))
        tk.Entry(f_frame, textvariable=self.dir_var, state="readonly").grid(
            row=0, column=1, sticky="we", padx=5
        )

        tk.Label(f_frame, text="File:").grid(row=1, column=0, sticky="e")
        last_name = Path(self.cfg.get("last_file")).name if self.cfg.get("last_file") else ""
        self.file_var = tk.StringVar(value=last_name)
        tk.Entry(f_frame, textvariable=self.file_var, state="readonly").grid(
            row=1, column=1, sticky="we", padx=5
        )
        tk.Button(f_frame, text="Browse...", command=self.on_browse).grid(row=1, column=2)
        f_frame.columnconfigure(1, weight=1)

        a_frame = tk.Frame(self.root)
        a_frame.pack(fill="x", **pad)
        tk.Button(
            a_frame,
            text="Generate Invoice Email",
            command=self.on_generate,
            width=25,
            bg="#e1f5fe",
        ).pack(side="left", padx=5)
        tk.Button(
            a_frame,
            text="Roll to Next Month",
            command=self.on_roll,
            width=25,
        ).pack(side="left", padx=5)

        self._panel_open = tk.BooleanVar(value=False)

        toggle_frame = tk.Frame(self.root)
        toggle_frame.pack(fill="x", padx=15, pady=(8, 0))
        self._toggle_btn = tk.Button(
            toggle_frame,
            text="▶ Invoice Information",
            anchor="w",
            relief="flat",
            cursor="hand2",
            command=self._toggle_panel,
        )
        self._toggle_btn.pack(fill="x")

        self._panel = tk.Frame(self.root, bd=1, relief="groove")

        tk.Label(self._panel, text="To:").grid(row=0, column=0, sticky="e", padx=(10, 2), pady=2)
        self.to_ent = tk.Entry(self._panel)
        self.to_ent.grid(row=0, column=1, sticky="we", pady=2, padx=(0, 10))

        tk.Label(self._panel, text="CC:").grid(row=1, column=0, sticky="e", padx=(10, 2), pady=2)
        self.cc_ent = tk.Entry(self._panel)
        self.cc_ent.grid(row=1, column=1, sticky="we", pady=2, padx=(0, 10))

        tk.Label(self._panel, text="Subject:").grid(row=2, column=0, sticky="e", padx=(10, 2), pady=2)
        self.subj_ent = tk.Entry(self._panel)
        self.subj_ent.grid(row=2, column=1, sticky="we", pady=2, padx=(0, 10))

        self.body_txt = scrolledtext.ScrolledText(self._panel, height=12, wrap="word")
        self.body_txt.grid(row=3, column=0, columnspan=2, sticky="nsew", pady=5, padx=10)
        self._panel.rowconfigure(3, weight=1)
        self._panel.columnconfigure(1, weight=1)

        b_frame = tk.Frame(self._panel)
        b_frame.grid(row=4, column=0, columnspan=2, sticky="we", padx=10, pady=(0, 8))
        tk.Button(b_frame, text="Copy Body", command=self.on_copy).pack(side="left", padx=5)
        tk.Button(
            b_frame,
            text="Open in Mailer (With CC)",
            bg="#c8e6c9",
            command=self.on_mailto,
        ).pack(side="left", padx=5)

        s_frame = tk.Frame(self.root)
        s_frame.pack(fill="x", padx=15, pady=2)

        self.rem_var = tk.BooleanVar(value=self.cfg.get_bool("remember_file", True))
        tk.Checkbutton(
            s_frame,
            text="Remember last file",
            variable=self.rem_var,
            command=self.update_prefs,
        ).pack(side="left")

        self.log_var = tk.BooleanVar(value=self.cfg.get_bool("enable_logging"))
        tk.Checkbutton(
            s_frame,
            text="Enable Logging",
            variable=self.log_var,
            command=self.update_prefs,
        ).pack(side="left", padx=20)

        self.status_var = tk.StringVar(value="Ready.")
        tk.Label(
            self.root,
            textvariable=self.status_var,
            bd=1,
            relief="sunken",
            anchor="w",
            pady=4,
        ).pack(fill="x", side="bottom")

    def _selected_path(self) -> Path:
        return Path(self.dir_var.get()) / self.file_var.get()

    def _toggle_panel(self) -> None:
        if self._panel_open.get():
            self._panel.pack_forget()
            self._toggle_btn.config(text="▶ Invoice Information")
            self._panel_open.set(False)
            self.root.geometry(WINDOW_COLLAPSED)
        else:
            self._panel.pack(fill="both", expand=True, padx=15, pady=(0, 5))
            self._toggle_btn.config(text="▼ Invoice Information")
            self._panel_open.set(True)
            self.root.geometry(WINDOW_EXPANDED)

    def _set_file(self, path: Path) -> None:
        path = path.expanduser().resolve()
        self.dir_var.set(str(path.parent))
        self.file_var.set(path.name)
        if self.rem_var.get():
            self.cfg.update(last_file=str(path), last_directory=str(path.parent))

    def _fill_invoice_fields(self, email: InvoiceEmail) -> None:
        self.to_ent.delete(0, tk.END)
        self.to_ent.insert(0, email.to)
        self.cc_ent.delete(0, tk.END)
        self.cc_ent.insert(0, email.cc)
        self.subj_ent.delete(0, tk.END)
        self.subj_ent.insert(0, email.subject)
        self.body_txt.delete("1.0", tk.END)
        self.body_txt.insert("1.0", email.body)

    def _panel_email(self) -> InvoiceEmail:
        return InvoiceEmail(
            to=self.to_ent.get(),
            cc=self.cc_ent.get(),
            subject=self.subj_ent.get(),
            body=self.body_txt.get("1.0", tk.END).strip(),
        )

    def _report_workbook_error(self, path: Path, exc: BaseException, context: str) -> None:
        if isinstance(exc, PermissionError) or is_permission_denied(exc):
            message = str(exc) if isinstance(exc, PermissionError) else workbook_locked_message(path)
            show_error(self.root, message)
            self.status_var.set("Permission denied - close Excel first.")
            return
        logging.error("%s: %s", context, exc)
        show_error(self.root, f"{context}:\n{exc}")
        self.status_var.set(f"{context} failed.")

    def update_prefs(self) -> None:
        self.cfg.update(
            remember_file=str(self.rem_var.get()),
            enable_logging=str(self.log_var.get()),
        )
        setup_logging(self.cfg)

    def on_browse(self) -> None:
        chosen = filedialog.askopenfilename(
            initialdir=self.dir_var.get() or str(Path.home()),
            filetypes=[("Excel", "*.xlsx")],
        )
        if chosen:
            self._set_file(Path(chosen))

    def on_generate(self) -> None:
        path = self._selected_path()
        if not path.exists():
            show_error(self.root, f"File not found:\n{path}")
            return
        try:
            self.status_var.set(f"Reading {path.name}...")
            email = build_email(parse_workbook(path))
            self._fill_invoice_fields(email)
            logging.info("Generated invoice for %s", path.name)
            if not self._panel_open.get():
                self._open_mailto(email)
                self.status_var.set("Opened mailer.")
            else:
                self.status_var.set("Invoice email generated successfully.")
        except Exception as e:
            self._report_workbook_error(path, e, "Failed to parse workbook")

    def on_roll(self) -> None:
        path = self._selected_path()
        if not path.exists():
            show_error(self.root, f"File not found:\n{path}")
            return
        try:
            sheet_month = read_sheet_month(path)
            next_start = next_month_start(sheet_month)
            current_start = month_start(datetime.now())

            if sheet_month < current_start:
                target = ask_catch_up_roll(
                    self.root, sheet_month, next_start, current_start
                )
                if target is None:
                    return
            else:
                if not messagebox.askyesno(
                    "Confirm", f"Create next-month roll from {path.name}?"
                ):
                    return
                target = next_start

            self.status_var.set(f"Rolling to {target.strftime('%B %Y')}...")
            new_path = roll_month(path, target)
            self._set_file(new_path)
            self.status_var.set(f"Rolled: {new_path.name}")
            messagebox.showinfo("Success", f"Created: {new_path.name}")
            logging.info("Rolled workbook to %s", new_path.name)
        except Exception as e:
            self._report_workbook_error(path, e, "Roll failed")

    def on_copy(self) -> None:
        self.root.clipboard_clear()
        self.root.clipboard_append(self.body_txt.get("1.0", tk.END).strip())
        self.status_var.set("Email body copied to clipboard.")

    def _open_mailto(self, email: InvoiceEmail) -> None:
        # CRLF is more reliable for Windows mailers; indent uses NBSP in the body.
        body = email.body.replace("\r\n", "\n").replace("\n", "\r\n")
        url = (
            f"mailto:{quote(email.to, safe='@')}"
            f"?cc={quote(email.cc, safe='@')}"
            f"&subject={quote(email.subject)}"
            f"&body={quote(body)}"
        )
        webbrowser.open(url)
        logging.info("Opened mailto link.")

    def on_mailto(self) -> None:
        self._open_mailto(self._panel_email())
        self.status_var.set("Opened mailer.")


if __name__ == "__main__":
    root = tk.Tk()
    if openpyxl is None:
        root.withdraw()
        messagebox.showerror(
            "Error",
            "Required library 'openpyxl' not found.\nRun: pip install openpyxl",
        )
        sys.exit(1)
    BillingAssistantGUI(root, sys.argv)
    root.mainloop()
